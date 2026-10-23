"""Exports des transcriptions (TXT, Markdown, JSON) et des comptes rendus (Markdown, DOCX)."""

from __future__ import annotations

import json
import re
from pathlib import Path

from teams_assistant.models import MeetingSession, Transcript, TranscriptSegment
from teams_assistant.summary.schema import EXTRACT, INTERPRETATION, SECTIONS, SummaryContent, SummaryItem


def timestamp(seconds: float) -> str:
    minutes, sec = divmod(int(seconds), 60)
    hours, minutes = divmod(minutes, 60)
    return f"{hours:02d}:{minutes:02d}:{sec:02d}"


def safe_filename(title: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", title).strip(" .")
    return cleaned[:80] or "reunion"


def _line(seg: TranscriptSegment) -> str:
    flag = " [à vérifier]" if seg.uncertain else ""
    speaker = f"{seg.speaker} : " if seg.speaker else ""
    return f"[{timestamp(seg.start_s)}] {speaker}{seg.text}{flag}"


def transcript_txt(session: MeetingSession, segments: list[TranscriptSegment]) -> str:
    header = [session.title, f"Durée : {timestamp(session.duration_s)}", ""]
    return "\n".join(header + [_line(s) for s in segments]) + "\n"


def transcript_md(session: MeetingSession, transcript: Transcript, segments: list[TranscriptSegment]) -> str:
    lines = [f"# Transcription — {session.title}", "",
             f"- Origine : {'transcription locale' if transcript.origin.value == 'local' else 'transcript Teams'}"
             f" ({transcript.engine} {transcript.model})",
             f"- Langue : {transcript.language}", "- Intervenants : sources audio, non identifiées", ""]
    lines += [f"- **[{timestamp(s.start_s)}]** #{s.seq} {s.speaker or ''} — {s.text}"
              + (" ⚠️ *à vérifier*" if s.uncertain else "") for s in segments]
    return "\n".join(lines) + "\n"


def transcript_json(session: MeetingSession, transcript: Transcript, segments: list[TranscriptSegment]) -> str:
    payload = {
        "session": {"id": session.id, "title": session.title, "duration_s": session.duration_s},
        "transcript": {"id": transcript.id, "origin": transcript.origin.value, "engine": transcript.engine,
                       "model": transcript.model, "language": transcript.language,
                       "warnings": transcript.warnings},
        "segments": [{"seq": s.seq, "start_s": s.start_s, "end_s": s.end_s, "source": s.source,
                      "speaker": s.speaker, "text": s.text, "original_text": s.original_text,
                      "confidence": s.confidence, "uncertain": s.uncertain} for s in segments],
    }
    return json.dumps(payload, ensure_ascii=False, indent=2)


def _item_md(item: SummaryItem) -> str:
    tags = []
    if item.kind == INTERPRETATION:
        tags.append("*interprétation*")
    elif item.kind == EXTRACT:
        tags.append("*extrait*")
    if item.owner:
        tags.append(f"responsable : {item.owner}")
    if item.due:
        tags.append(f"échéance : {item.due}")
    if item.speaker:
        tags.append(f"source : {item.speaker}")
    if item.refs:
        tags.append("segments " + ", ".join(f"#{r}" for r in item.refs))
    return f"- {item.text}" + (f" ({' · '.join(tags)})" if tags else "")


def summary_md(content: SummaryContent, validated: bool) -> str:
    lines = [f"# {content.title}", ""]
    if content.date or content.duration:
        lines += [f"**Date** : {content.date or 'inconnue'} · **Durée** : {content.duration or 'inconnue'}", ""]
    lines.append(f"**Statut** : {'validé par l’utilisateur' if validated else 'BROUILLON — non validé'}")
    if content.generator_notice:
        lines += ["", f"> {content.generator_notice}"]
    lines += ["", "## Participants", ""]
    lines += [f"- {p.name} — identité {p.status} ({p.origin})" for p in content.participants] or ["- Inconnus"]
    for key, label in SECTIONS:
        lines += ["", f"## {label}", ""]
        items = content.sections.get(key, [])
        lines += [_item_md(i) for i in items] if items else ["- Aucun élément identifié."]
    return "\n".join(lines) + "\n"


def summary_docx(content: SummaryContent, validated: bool, path: Path) -> Path:
    from docx import Document  # python-docx

    doc = Document()
    doc.add_heading(content.title, level=0)
    doc.add_paragraph(f"Date : {content.date or 'inconnue'} — Durée : {content.duration or 'inconnue'}")
    doc.add_paragraph("Statut : " + ("validé par l'utilisateur" if validated else "BROUILLON — non validé"))
    if content.generator_notice:
        doc.add_paragraph().add_run(content.generator_notice).italic = True
    doc.add_heading("Participants", level=1)
    for p in content.participants:
        doc.add_paragraph(f"{p.name} — identité {p.status} ({p.origin})", style="List Bullet")
    for key, label in SECTIONS:
        doc.add_heading(label, level=1)
        items = content.sections.get(key, [])
        for item in items or [SummaryItem("Aucun élément identifié.")]:
            doc.add_paragraph(_item_md(item)[2:], style="List Bullet")
    doc.save(str(path))
    return path


def write_text(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path
